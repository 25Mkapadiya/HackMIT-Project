#!/usr/bin/env python3
"""
Build public/data/power-plants.json from EIA-860 (Schedule 2 Plant + Schedule
3.1 Generator, Operable), for every US state — not filtered to one state.
Source of truth: https://www.eia.gov/electricity/data/eia860/

This is the free, nationwide, plant-level substitute for the CEII-restricted
substation-headroom layer: real substation interconnection capacity isn't
public (18 CFR 388.113), so this app uses generation capacity + grid voltage
near a site as an *estimated* proxy for transmission density instead. Every
record this script emits should be treated as "nearby transmission
infrastructure", not "available headroom".

Usage:
    python3 scripts/build_power_layer.py <path-to-extracted-eia860-dir>

Expects the extracted EIA-860 annual zip to contain:
    2___Plant_Y####.xlsx
    3_1_Generator_Y####.xlsx  (sheet "Operable")
"""
import glob
import json
import os
import sys

import openpyxl


def find_one(directory, pattern):
    matches = glob.glob(os.path.join(directory, pattern))
    if not matches:
        raise SystemExit(f"no file matching {pattern} in {directory}")
    return matches[0]


def load_rows(path, sheet=None, header_row=2):
    """EIA-860 sheets have a title row, then a header row, then data."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
    rows = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rows.append(dict(zip(header, r)))
    return rows


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    src_dir = sys.argv[1]

    plant_path = find_one(src_dir, "2___Plant_Y*.xlsx")
    gen_path = find_one(src_dir, "3_1_Generator_Y*.xlsx")

    print(f"reading {plant_path}")
    plants = load_rows(plant_path)
    all_plants = {p["Plant Code"]: p for p in plants if p.get("State")}
    print(f"  {len(all_plants)} plants across all states")

    print(f"reading {gen_path} (Operable)")
    gens = load_rows(gen_path, sheet="Operable")

    capacity_by_plant = {}
    tech_by_plant = {}
    for g in gens:
        code = g.get("Plant Code")
        if code not in all_plants:
            continue
        mw = g.get("Nameplate Capacity (MW)") or 0
        try:
            mw = float(mw)
        except (TypeError, ValueError):
            mw = 0
        capacity_by_plant[code] = capacity_by_plant.get(code, 0.0) + mw
        tech_by_plant.setdefault(code, set()).add(g.get("Technology") or "unknown")

    features = []
    skipped_no_coords = 0
    for code, p in all_plants.items():
        lat, lon = p.get("Latitude"), p.get("Longitude")
        if lat in (None, "", " ") or lon in (None, "", " "):
            skipped_no_coords += 1
            continue
        mw = round(capacity_by_plant.get(code, 0.0), 1)
        voltages = [p.get("Grid Voltage (kV)"), p.get("Grid Voltage 2 (kV)"), p.get("Grid Voltage 3 (kV)")]
        voltages = [v for v in voltages if isinstance(v, (int, float)) and v > 0]
        features.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [round(float(lon), 5), round(float(lat), 5)]},
            "properties": {
                "plant_code": code,
                "name": p.get("Plant Name"),
                "state": p.get("State"),
                "county": p.get("County"),
                "utility": p.get("Utility Name"),
                "balancing_authority": p.get("Balancing Authority Code"),
                "nameplate_mw": mw,
                "max_grid_voltage_kv": max(voltages) if voltages else None,
                "technologies": sorted(t for t in tech_by_plant.get(code, []) if t),
            },
        })

    print(f"  {len(features)} plants with coordinates ({skipped_no_coords} skipped, no lat/lon)")

    out = {
        "type": "FeatureCollection",
        "metadata": {
            "source": "EIA-860 (2025 annual), Schedule 2 Plant + Schedule 3.1 Generator (Operable)",
            "source_url": "https://www.eia.gov/electricity/data/eia860/",
            "generated_by": "scripts/build_power_layer.py",
            "estimate": True,
            "estimate_reason": "Proxy for transmission/substation headroom, which is CEII-restricted "
                                "(18 CFR 388.113) and not publicly available.",
        },
        "features": features,
    }

    out_path = os.path.join(os.path.dirname(__file__), "..", "public", "data", "power-plants.json")
    with open(out_path, "w") as f:
        json.dump(out, f, separators=(",", ":"))
    print(f"wrote {out_path} ({os.path.getsize(out_path)} bytes)")


if __name__ == "__main__":
    main()
